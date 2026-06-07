import os, sys, re, json, glob, subprocess
from copy import deepcopy
from pathlib import Path
from dataclasses import asdict
from models import SubjectRecord, SemesterRecord, StudentRecord

try:
    import openpyxl
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    HAS_OPENPYXL = True
except ImportError: HAS_OPENPYXL = False

try: import pandas as pd; HAS_PANDAS = True
except ImportError: HAS_PANDAS = False

FORMAT_CONFIGS = {
    "aktu_f1": {"name": "AKTU OneView Format 1", "detect_keywords": ["OneView", "SGPA"], "header_regex": {"institute_code": r'Institute Code\s*\n\s*:\s*\((\d+)\s*\)', "course": r'\(0\d\)\s*(B\.TECH|M\.TECH|MBA|MCA|B\.PHARM)', "roll_number": r'RollNo\s*:\s*(\S+)', "enrollment_number": r'EnrollmentNo\s*:\s*(\S+)', "student_name": r'Name\s*:\s*([A-Z][A-Z\s]+?)(?:\s{2,}|Hindi Name|Father)', "father_name": r"Father'?s?\s*Name\s*:\s*([A-Z][A-Z\s]+?)(?:\s{2,}|Gender|$)", "gender": r'Gender\s*:\s*([MF])\b'}, "table_delegate": "parse_subjects_f1"},
    "aktu_f2": {"name": "AKTU Tabulation Format 2", "detect_keywords": ["Tabulation", "NAME OF INSTITUTE", "CP("], "header_regex": {}, "table_delegate": "parse_student_block_f2"}
}

def detect_format(text: str) -> str:
    txt_upper = text.upper()
    if "ONEVIEW" in txt_upper or "DATE OF DECLARATION" in txt_upper: return "aktu_f1"
    elif "TABULATION" in txt_upper or "NAME OF INSTITUTE" in txt_upper or "CP(" in txt_upper: return "aktu_f2"
    return "unknown"

def validate_and_normalise(raw_dict: dict, source: str) -> dict:
    REQUIRED = ["student_name", "roll_number", "semesters"]
    for fld in REQUIRED:
        if not raw_dict.get(fld): raise ValueError(f"[{source}] Validation Failed: Missing required field '{fld}'.")
    for sem in raw_dict.get("semesters", []):
        if sem.get("sgpa") is None: sem["sgpa"] = None
    return raw_dict

def extract_header_via_config(text: str, config_key: str) -> dict:
    cfg = FORMAT_CONFIGS.get(config_key, {}).get("header_regex", {})
    result = {}
    for key, pattern in cfg.items():
        m = re.search(pattern, text, re.IGNORECASE)
        result[key] = m.group(1).strip() if m else None
    return result

def _f(val):
    if val is None: return None
    try: return float(str(val).strip().rstrip("*#"))
    except: return None

def _pm(raw):
    if not raw or str(raw).strip() in ("--",""): return None, False
    raw = str(raw).strip()
    return _f(raw), raw.endswith("*")

def _int_re(text, pat):
    m = re.search(pat, text, re.IGNORECASE); return int(m.group(1)) if m else None
def _flt_re(text, pat):
    m = re.search(pat, text, re.IGNORECASE); return float(m.group(1)) if m else None
def _str_re(text, pat):
    m = re.search(pat, text, re.IGNORECASE); return m.group(1).strip() if m else None

def extract_pdf_text_f1(pdf_path):
    try:
        import pdfplumber; pages = []
        with pdfplumber.open(pdf_path) as pdf:
            for page in pdf.pages:
                t = page.extract_text(); 
                if t: pages.append(t)
        return "\n".join(pages)
    except ImportError:
        from pypdf import PdfReader; r = PdfReader(pdf_path)
        return "\n".join(p.extract_text() or "" for p in r.pages)

SUBJECT_RE_F1 = re.compile(r'^([A-Z]{2,5}\d{3,4}[A-Z]?)\s+(.*?)\s+(Theory|Practical|CA)\s+(\d{1,3}|--)\s+(\d{1,3}\*?|--)\s*(?:(\d{1,3}\*?|--)\s*)?([A-Z][+#]?|--)?\s*$', re.IGNORECASE)

def parse_student_header_f1(text, source_file=None):
    extracted = extract_header_via_config(text, "aktu_f1")
    def find(pat, default=None):
        m = re.search(pat, text, re.IGNORECASE); return m.group(1).strip() if m else default
    inst_name = None
    inst_m = re.search(r'Institute Code\s*[&\n]+\s*Name\s*:\s*\(\d+\s*\)\s*(.+?)(?:\n|Course)', text, re.IGNORECASE)
    if inst_m: inst_name = inst_m.group(1).strip()
    branch = None
    bm = re.search(r'Branch\s*Code[^\(]*\(\d{2,3}\)\s*([\s\S]{5,100}?)(?:\n|RollNo|Roll|Enrollment)', text, re.IGNORECASE)
    if bm:
        raw_b = bm.group(1); raw_b = re.sub(r'\(0\d\)\s*[A-Z\.]+', ' ', raw_b, flags=re.IGNORECASE); raw_b = re.sub(r'\bName\b\s*:?', '', raw_b, flags=re.IGNORECASE); raw_b = re.sub(r'\s*:[\s\S]*?:\s*', ' ', raw_b)
        branch = re.sub(r'\s+', ' ', raw_b).strip()
    if not branch: branch = find(r'Branch Code\s*[&\n]+\s*Name\s*:\s*\(\d+\)\s*(.+?)(?:\n|RollNo|Roll)')
    name = extracted.get("student_name")
    if not name: name = find(r'Name\s*:\s*([A-Z ]{5,40})')
    return StudentRecord(source_file=source_file, institute_code=extracted.get("institute_code"), institute_name=inst_name, course=extracted.get("course"), branch=branch, roll_number=extracted.get("roll_number"), enrollment_number=extracted.get("enrollment_number"), student_name=name, father_name=extracted.get("father_name"), gender=extracted.get("gender"))

def parse_subjects_f1(block_lines):
    subjects = []
    for line in block_lines:
        line = line.strip(); 
        if not line: continue
        m = SUBJECT_RE_F1.match(line); 
        if not m: continue
        code, name_, stype, internal, external, back, grade = m.groups()
        name_clean = name_.strip().rstrip("*#").strip()
        int_val = _f(internal) if internal and internal != "--" else None
        ext_val, _ = _pm(external); back_val, _ = _pm(back)
        grade_val = grade.strip() if grade and grade.strip() not in ("--","") else None
        total = None
        if int_val is not None and ext_val is not None: total = int_val + ext_val
        elif int_val is not None: total = int_val
        subjects.append(SubjectRecord(subject_code=code.strip(), subject_name=name_clean, subject_type=stype.strip(), internal_marks=int_val, external_marks=ext_val, total_marks=total, back_paper_marks=back_val, grade=grade_val))
    return subjects

def split_semesters_f1(text):
    sem_re = re.compile(r'Semester\s*:\s*(\d+)\s+Even/Odd\s*:\s*(Odd|Even)', re.IGNORECASE)
    matches = list(sem_re.finditer(text))
    if not matches: return []
    blocks = []
    for i, m in enumerate(matches):
        start = m.start(); end = matches[i+1].start() if i+1 < len(matches) else len(text); bt = text[start:end]
        best_session = None
        for sm in re.finditer(r'Session\s*:\s*([\w\-]+\(\w+\))', text[:start], re.IGNORECASE): best_session = sm.group(1).strip()
        blocks.append({"semester_number": int(m.group(1)), "even_odd": m.group(2), "session": best_session, "total_subjects": _int_re(bt, r'Total Subjects\s*:\s*(\d+)'), "theory_subjects": _int_re(bt, r'Theory Subjects\s*:\s*(\d+)'), "practical_subjects":_int_re(bt, r'Practical Subjects\s*:\s*(\d+)'), "total_marks_obtained":_flt_re(bt, r'Total Marks Obt\.\s*:\s*(\d+(?:\.\d+)?)'), "result_status": _str_re(bt, r'Result Status\s*:\s*([A-Z]+(?:\(\s*\d+\s*\))?)'), "sgpa": _flt_re(bt, r'SGPA\s*:\s*(\d+(?:\.\d+)?)'), "date_of_declaration":_str_re(bt, r'Date of Declaration\s*:\s*(\d{2}/\d{2}/\d{2,4})'), "lines": bt.splitlines()})
    return blocks

def find_subject_table_start(lines):
    for i, line in enumerate(lines):
        if re.search(r'\bCode\b.*\bName\b.*\bType\b.*\bInternal\b', line, re.IGNORECASE): return i + 1
    return 0

def parse_aktu_pdf_f1(pdf_path, raw_text=None):
    raw = raw_text if raw_text else extract_pdf_text_f1(pdf_path)
    student = parse_student_header_f1(raw, source_file=os.path.basename(pdf_path))
    sem_map = {}
    for blk in split_semesters_f1(raw):
        start = find_subject_table_start(blk["lines"])
        subjects = parse_subjects_f1(blk["lines"][start:])
        sem = SemesterRecord(semester_number=blk["semester_number"], even_odd=blk["even_odd"], session=blk["session"], total_subjects=blk["total_subjects"], theory_subjects=blk["theory_subjects"], practical_subjects=blk["practical_subjects"], total_marks_obtained=blk["total_marks_obtained"], result_status=blk["result_status"], sgpa=blk["sgpa"], date_of_declaration=blk["date_of_declaration"], subjects=subjects)
        sn = sem.semester_number
        if sn not in sem_map: sem_map[sn] = sem
        else:
            ex = sem_map[sn]; ns, es = sem.total_marks_obtained or 0, ex.total_marks_obtained or 0
            if ns > es or (ns == es and "PASS" in (sem.result_status or "") and "PASS" not in (ex.result_status or "")): sem_map[sn] = sem
    student.semesters = [sem_map[k] for k in sorted(sem_map)]
    return student

def process_batch_f1(input_dir, output_dir, progress_cb=None):
    input_dir  = os.path.abspath(input_dir); output_dir = os.path.abspath(output_dir)
    os.makedirs(output_dir, exist_ok=True)
    pdf_files = sorted(set(glob.glob(os.path.join(input_dir, "*.pdf")) + glob.glob(os.path.join(input_dir, "**", "*.pdf"), recursive=True)))
    results = {}
    for idx, pdf_path in enumerate(pdf_files):
        short = os.path.relpath(pdf_path, input_dir)
        try:
            raw_text = extract_pdf_text_f1(pdf_path); fmt = detect_format(raw_text)
            if fmt == "unknown": raise ValueError("PDF format unrecognized. Expected AKTU OneView.")
            record = parse_aktu_pdf_f1(pdf_path, raw_text=raw_text)
            record_dict = validate_and_normalise(asdict(record), short)
            stem = record_dict.get("roll_number") or Path(pdf_path).stem
            out_path = os.path.join(output_dir, re.sub(r'[^\w]','_',stem) + ".json")
            with open(out_path, "w", encoding="utf-8") as f: json.dump(record_dict, f, indent=2, ensure_ascii=False)
            results[short] = {"status":"ok", "student":record_dict.get("student_name"), "roll":record_dict.get("roll_number"), "semesters":len(record_dict.get("semesters", []))}
        except Exception as e: results[short] = {"status":f"error: {e}"}
        if progress_cb: progress_cb(idx+1, len(pdf_files), short)
    return results

SUBJECT_META_F2 = {"BCE501":("Structural Analysis-II","Theory"), "BCE502":("Design of Steel Structures","Theory"), "BCE503":("Design of Concrete Structures","Theory"), "BCE051":("Elective-I","Theory"), "BCE055":("Elective-II","Theory"), "BCE551":("Structural Analysis-II Lab","Practical"), "BCE552":("Design of Steel Structures Lab","Practical"), "BCE553":("Design of Concrete Structures Lab","Practical"), "BCE554":("Minor Project / Internship Assessment","CA"), "BNC501":("Constitution of India","CA")}

def extract_pdf_text_f2(pdf_path):
    try:
        import pdfplumber; pages = []
        with pdfplumber.open(pdf_path) as pdf:
            for page in pdf.pages:
                try: t = page.extract_text(layout=True)
                except TypeError: t = page.extract_text()
                if t: pages.append(t)
        if pages: return "\n".join(pages)
    except ImportError: pass
    r = subprocess.run(["pdftotext","-layout",pdf_path,"-"], capture_output=True, text=True)
    if r.returncode == 0 and r.stdout.strip(): return r.stdout
    raise RuntimeError("Cannot extract PDF text — install pdfplumber")

def _num_toks(line, count=10):
    result = []
    for t in re.findall(r'\S+', line):
        if re.match(r'^\d+$|^--$', t):
            result.append(t)
            if len(result) == count: break
    return result

def _name_from_line(line):
    parts = []
    for t in re.findall(r'\S+', line):
        if re.match(r'^\d+$|^--$', t): break
        parts.append(t)
    return " ".join(parts)

def parse_format2_pdf(pdf_path):
    raw = extract_pdf_text_f2(pdf_path); fmt = detect_format(raw)
    if fmt != "aktu_f2": raise ValueError("PDF format unrecognized. Expected AKTU Tabulation Register.")
    lines = raw.split("\n")
    starts = [i for i, line in enumerate(lines) if re.search(r'\d{1,3}\s+\d{13}', line) and re.search(r'[A-Z]{2,5}\d{3}', line)]
    students = []; WORD_NUM = {"FIRST":1,"SECOND":2,"THIRD":3,"FOURTH":4,"FIFTH":5,"SIXTH":6,"SEVENTH":7,"EIGHTH":8}
    current_doc = {"institute_code": None, "institute_name": None, "branch_name": None, "course": "B.TECH", "sem_word": None, "exam_year": None, "semester_number": None}
    expecting_inst = False; expecting_br = False; first_doc = None
    for i, line in enumerate(lines):
        m_sem = re.search(r'(BACHELOR|MASTER)\s+OF\s+.*?(\w+)\s+SEMESTER\s+EXAM\.?\s*([\d\-\/]+)', line, re.IGNORECASE)
        if m_sem: current_doc["course"] = m_sem.group(1).upper() + ".TECH" if "TECHNOLOGY" in line.upper() else m_sem.group(1).upper(); current_doc["sem_word"] = m_sem.group(2).upper(); current_doc["exam_year"] = m_sem.group(3).strip(); current_doc["semester_number"] = WORD_NUM.get(current_doc["sem_word"], None)
        if "NAME OF INSTITUTE" in line.upper():
            m = re.search(r'NAME OF INSTITUTE\s*:\s*(\d{3})\s+([^\s]{2,}.*?)(?:\s{3,}|$)', line, re.IGNORECASE)
            if m: current_doc["institute_code"] = m.group(1).strip(); current_doc["institute_name"] = m.group(2).strip()
            else: expecting_inst = True
        if "NAME OF BRANCH" in line.upper():
            m = re.search(r'NAME OF BRANCH\s*:\s*(\d{2,3})\s+([^\s]{2,}.*?)(?:\s{3,}|$)', line, re.IGNORECASE)
            if m: current_doc["branch_name"] = m.group(2).strip()
            else: expecting_br = True
        if expecting_inst:
            m = re.search(r'(?:^|\s{2,})(\d{3})\s+([A-Z][A-Z\s&,\.\-]{10,})(?:\s{2,}|$)', line)
            if m: current_doc["institute_code"] = m.group(1).strip(); current_doc["institute_name"] = m.group(2).strip(); expecting_inst = False
        if expecting_br:
            m = re.search(r'(?:^|\s{2,})(\d{2,3})\s+([A-Z][A-Z\s\.\-\(\)]{4,})(?:\s{2,}|$)', line)
            if m: current_doc["branch_name"] = m.group(2).strip(); expecting_br = False
        if not first_doc and current_doc["institute_code"] and current_doc["branch_name"]: first_doc = current_doc.copy()
        if i in starts:
            student = parse_student_block_f2(lines, i, current_doc.copy())
            if student: students.append(student)
    rep_doc = first_doc if first_doc else current_doc.copy()
    if students:
        unique_branches = list(set(s.get("_branch") for s in students if s.get("_branch")))
        if len(unique_branches) > 1: rep_doc["branch_name"] = f"Multiple ({len(unique_branches)} branches)"
        elif len(unique_branches) == 1: rep_doc["branch_name"] = unique_branches[0]
    return rep_doc, students

def parse_student_block_f2(lines, start_idx, doc_header):
    try:
        l0 = lines[start_idx]; m_r = re.search(r'(\d{13})', l0)
        if not m_r: return None
        roll = m_r.group(1); codes = [c for c in re.findall(r'[A-Z]{2,5}\d{3}', l0) if len(c)>=5]; n = len(codes)
        if n == 0: return None
        l1 = lines[start_idx+1]; ext_tok = _num_toks(l1, n); m_res = re.search(r'(\d+)\s+CP\(\s*(\d+)\s*\)\s+S\.G\.P\.A\.\s*:\s*([\d.]+)', l1)
        grand_total = int(m_res.group(1)) if m_res else None; cp_count = int(m_res.group(2)) if m_res else 0; sgpa = float(m_res.group(3)) if m_res else None
        l2 = lines[start_idx+2]; student_name = _name_from_line(l2); int_tok = _num_toks(l2, n)
        l3 = lines[start_idx+3]; father_name  = _name_from_line(l3); tot_tok = _num_toks(l3, n)
        l4 = lines[start_idx+4]; grades = re.findall(r'[A-F][+#]?', l4)
        back_paper_codes = []
        if start_idx+6 < len(lines):
            l6 = lines[start_idx+6].strip()
            if (re.search(r'[A-Z]{2,5}\d{3}', l6) and '---' not in l6 and 'PAGE' not in l6 and 'BACHELOR' not in l6 and 'SNO' not in l6): back_paper_codes = re.findall(r'[A-Z]{2,5}\d{3}', l6)
        sem_num = doc_header.get("semester_number"); subjects = []
        for i, code in enumerate(codes):
            meta_name, meta_type = SUBJECT_META_F2.get(code, (code,"Theory"))
            raw_ext = ext_tok[i] if i<len(ext_tok) else "--"; raw_int = int_tok[i] if i<len(int_tok) else "--"; raw_tot = tot_tok[i] if i<len(tot_tok) else "--"
            ext = None if raw_ext=="--" else int(raw_ext); int_ = None if raw_int=="--" else int(raw_int); tot = None if raw_tot=="--" else int(raw_tot); grade = grades[i] if i<len(grades) else None
            subjects.append({"subject_code":code,"subject_name":meta_name,"subject_type":meta_type,"internal_marks":int_,"external_marks":ext,"total_marks":tot,"back_paper_marks":None,"grade":grade,"has_back_paper":code in back_paper_codes})
        even_odd = "Odd" if sem_num and sem_num%2==1 else "Even"
        return {"roll_number":roll,"student_name":student_name,"father_name":father_name,"semester_number":sem_num,"even_odd":even_odd,"session":doc_header.get("exam_year"),"sgpa":sgpa,"total_marks_obtained":grand_total,"result_status":f"CP( {cp_count})","back_paper_codes":back_paper_codes,"date_of_declaration":None,"total_subjects":n,"theory_subjects":sum(1 for s in subjects if s["subject_type"]=="Theory"),"practical_subjects":sum(1 for s in subjects if s["subject_type"]=="Practical"),"subjects":subjects,"_institute_code":doc_header.get("institute_code"),"_institute_name":doc_header.get("institute_name"),"_branch":doc_header.get("branch_name"), "_course":"B.TECH"}
    except Exception: return None

def run_format2_update(pdf_path, json_dir, dry_run=False, progress_cb=None):
    doc, students = parse_format2_pdf(pdf_path); existing = {}
    for fpath in glob.glob(os.path.join(json_dir,"*.json")):
        if os.path.basename(fpath).startswith("_"): continue
        try:
            with open(fpath,encoding="utf-8") as f: data=json.load(f)
            roll=data.get("roll_number")
            if roll: existing[roll]=(fpath,data)
        except: pass
    counters={"updated":0,"created":0,"skipped":0,"conflict":0}
    for idx,parsed in enumerate(students):
        parsed["_source_file"]=os.path.basename(pdf_path)
        roll=parsed["roll_number"]; sem_num=parsed["semester_number"]
        sem_rec={"semester_number":parsed["semester_number"],"even_odd":parsed["even_odd"],"session":parsed["session"],"total_subjects":parsed["total_subjects"],"theory_subjects":parsed["theory_subjects"],"practical_subjects":parsed["practical_subjects"],"total_marks_obtained":parsed["total_marks_obtained"],"result_status":parsed["result_status"],"sgpa":parsed["sgpa"],"date_of_declaration":parsed["date_of_declaration"],"subjects":[{k:v for k,v in s.items() if k!="has_back_paper"} for s in parsed["subjects"]]}
        if roll in existing:
            fpath,data=existing[roll]; data=deepcopy(data)
            if parsed.get("_branch") and data.get("branch") != parsed.get("_branch"): data["branch"] = parsed.get("_branch")
            already=[s for s in data.get("semesters",[]) if s["semester_number"]==sem_num]
            if already: result = "skipped" if already[0].get("sgpa")==parsed["sgpa"] else "conflict"
            else:
                data["semesters"].append(sem_rec); data["semesters"].sort(key=lambda s:s["semester_number"])
                data = validate_and_normalise(data, fpath)
                if not dry_run:
                    with open(fpath,"w",encoding="utf-8") as f: json.dump(data,f,indent=2,ensure_ascii=False)
                result="updated"
        else:
            new_data={"source_file":parsed.get("_source_file",""),"institute_code":parsed.get("_institute_code"),"institute_name":parsed.get("_institute_name"),"course":parsed.get("_course","B.TECH"),"branch":parsed.get("_branch"),"roll_number":roll,"enrollment_number":None,"student_name":parsed["student_name"],"father_name":parsed["father_name"],"gender":None,"semesters":[sem_rec]}
            new_data = validate_and_normalise(new_data, parsed.get("_source_file",""))
            out_path=os.path.join(json_dir,f"{roll}.json")
            if not dry_run:
                os.makedirs(json_dir,exist_ok=True)
                with open(out_path,"w",encoding="utf-8") as f: json.dump(new_data,f,indent=2,ensure_ascii=False)
            result="created"
        counters[result]+=1
        if progress_cb: progress_cb(idx+1,len(students),f"{parsed['student_name']} → {result}")
    return doc, students, counters

class DataExporter:
    @staticmethod
    def to_csv(data, filepath):
        if not data or not HAS_PANDAS: return
        pd.DataFrame(data).to_csv(filepath, index=False)
    @staticmethod
    def to_excel_simple(data, filepath):
        if not data: return
        if HAS_PANDAS: pd.DataFrame(data).to_excel(filepath, index=False)
        elif HAS_OPENPYXL:
            wb=Workbook(); ws=wb.active
            if data: ws.append(list(data[0].keys()))
            for row in data: ws.append(list(row.values()))
            wb.save(filepath)
    @staticmethod
    def to_pdf(data, filepath, title="University Report"):
        if not data: return
        try:
            from fpdf import FPDF
            if HAS_PANDAS: df=pd.DataFrame(data)
            else:
                class _DF:
                    def __init__(self,d): self.columns=list(d[0].keys()); self._d=d
                    def iterrows(self): return enumerate(self._d)
                df=_DF(data)
            pdf=FPDF(); pdf.add_page(); pdf.set_font("Arial",'B',14)
            pdf.cell(0,10,txt=title,ln=True,align='C'); pdf.ln(5)
            cols=df.columns.tolist() if hasattr(df,'columns') else list(data[0].keys())
            pdf.set_font("Arial",'B',9); col_width=190/len(cols) if cols else 190
            for col in cols: pdf.cell(col_width,10,str(col).replace('_',' ').title()[:15],border=1,align='C')
            pdf.ln(); pdf.set_font("Arial",'',8)
            for _,row in (df.iterrows() if hasattr(df,'iterrows') else enumerate(data)):
                row_data=row if isinstance(row,dict) else row
                for col in cols:
                    val=str(row_data.get(col,'') if isinstance(row_data,dict) else getattr(row_data,col,'')).encode('latin-1','replace').decode('latin-1')
                    pdf.cell(col_width,10,val[:20],border=1)
                pdf.ln()
            pdf.output(filepath)
        except ImportError: raise ImportError("Please install fpdf2: pip install fpdf2")
    @staticmethod
    def to_excel_master(qlib, out_path):
        if not HAS_OPENPYXL: raise ImportError("Install openpyxl: pip install openpyxl")
        sql="""SELECT RANK() OVER (ORDER BY st.cgpa DESC) as rank, st.name,st.roll_number,st.enrollment_number,st.institute_name,st.branch,st.course, st.cgpa,st.failed_subjects,sm.semester_number,sm.session,sm.sgpa,sm.result_status, sub.subject_code,sub.subject_name,sub.subject_type, sub.internal_marks,sub.external_marks,sub.total_marks,sub.back_paper_marks,sub.grade FROM students st LEFT JOIN semesters sm ON st.id=sm.student_id LEFT JOIN subjects sub ON sm.id=sub.semester_id ORDER BY st.cgpa DESC,st.roll_number,sm.semester_number,sub.subject_code"""
        rows=qlib._run(sql)
        if not rows: raise ValueError("No data in database to export.")
        def _fill(c): return PatternFill("solid",fgColor=c)
        def _font(bold=False,color="000000",size=10): return Font(name="Calibri",bold=bold,color=color,size=size)
        def _align(h="left"): return Alignment(horizontal=h,vertical="center")
        def _bdr(): s=Side(style="thin",color="CCCCCC"); return Border(left=s,right=s,top=s,bottom=s)
        wb=Workbook(); wb.remove(wb.active); ws=wb.create_sheet("Master Records")
        headers=["Rank","Student Name","Roll Number","Enrollment No","Institute","Branch","Course","CGPA","Fails","Sem","Session","SGPA","Status","Code","Subject Name","Type","Internal","External","Total","Back Paper","Grade"]
        col_widths={"A":6,"B":28,"C":16,"D":18,"E":36,"F":22,"G":10,"H":8,"I":8,"J":5,"K":14,"L":8,"M":14,"N":10,"O":38,"P":10,"Q":7,"R":7,"S":7,"T":12,"U":7}
        GRADE_CLR={"A+":"27AE60","A":"2ECC71","B+":"3498DB","B":"5DADE2","C":"F39C12","D":"E67E22","E":"E74C3C","E#":"C0392B","F":"922B21"}
        CENTER_COLS={1,8,9,10,12,17,18,19,20,21}
        for i,h in enumerate(headers,1):
            c=ws.cell(row=1,column=i,value=h)
            c.font=_font(bold=True,color="FFFFFF"); c.fill=_fill("1F3864"); c.alignment=_align("center"); c.border=_bdr()
        row_idx=2
        for r in rows:
            grade=r.get("grade") or ""; cgpa=r.get("cgpa")
            if grade in ("F","E","E#"): bg="FADBD8"
            elif r.get("back_paper_marks"): bg="FEF9E7"
            elif row_idx%2==0: bg="F2F2F2"
            else: bg="FFFFFF"
            vals=[r.get("rank"),r.get("name"),r.get("roll_number"),r.get("enrollment_number"), r.get("institute_name"),r.get("branch"),r.get("course"),cgpa,r.get("failed_subjects"), r.get("semester_number"),r.get("session"),r.get("sgpa"),r.get("result_status"), r.get("subject_code"),r.get("subject_name"),r.get("subject_type"), r.get("internal_marks"),r.get("external_marks"),r.get("total_marks"), r.get("back_paper_marks"),grade]
            for col,val in enumerate(vals,1):
                c=ws.cell(row=row_idx,column=col,value=val)
                c.fill=_fill(bg); c.alignment=_align("center" if col in CENTER_COLS else "left"); c.border=_bdr()
                if col==21 and grade: c.font=_font(bold=True,color=GRADE_CLR.get(grade,"000000"))
                elif col==8 and isinstance(cgpa,float): c.font=_font(bold=True,color="27AE60" if cgpa>=8.0 else ("3498DB" if cgpa>=6.5 else "E67E22"))
                else: c.font=_font()
            row_idx+=1
        for ltr,w in col_widths.items(): ws.column_dimensions[ltr].width=w
        ws.freeze_panes="A2"; ws.auto_filter.ref=f"A1:U{row_idx-1}"; wb.save(out_path)